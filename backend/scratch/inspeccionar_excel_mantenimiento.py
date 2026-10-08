"""Inspeccionar Excel CI-FR- INVENTARIO MANTENIMIENTO 2026.xlsx"""
import os
import openpyxl

ruta = r"c:\gs-viaticos\backend\data\migracion_inventario\CI-FR- INVENTARIO MANTENIMIENTO 2026.xlsx"
wb = openpyxl.load_workbook(ruta, data_only=True)
print("Hojas en el Excel:", wb.sheetnames)

for sheetname in wb.sheetnames:
    ws = wb[sheetname]
    print(f"\n--- Hoja: {sheetname} ---")
    print(f"Max row: {ws.max_row}, Max column: {ws.max_column}")
    # Primeras 5 filas
    for r in range(1, min(6, ws.max_row + 1)):
        vals = [ws.cell(r, c).value for c in range(1, min(12, ws.max_column + 1))]
        if any(v is not None for v in vals):
            print(f"  Fila {r}: {vals}")
