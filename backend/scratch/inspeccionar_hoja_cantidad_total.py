"""Inspeccionar hoja CANTIDAD TOTAL del Excel"""
import openpyxl

ruta = r"c:\gs-viaticos\backend\data\migracion_inventario\CI-FR- INVENTARIO MANTENIMIENTO 2026.xlsx"
wb = openpyxl.load_workbook(ruta, data_only=True)
ws = wb["CANTIDAD TOTAL"]

print(f"Max row: {ws.max_row}, Max column: {ws.max_column}")
for r in range(1, 15):
    vals = [ws.cell(r, c).value for c in range(1, min(15, ws.max_column + 1))]
    print(f"Fila {r}: {vals}")
