"""Inspeccionar hoja INVENTARIO GENERAL del Excel"""
import openpyxl

ruta = r"c:\gs-viaticos\backend\data\migracion_inventario\CI-FR- INVENTARIO MANTENIMIENTO 2026.xlsx"
wb = openpyxl.load_workbook(ruta, data_only=True)
ws = wb["INVENTARIO GENERAL"]

print(f"Rango: {ws.dimensions}, Max row: {ws.max_row}, Max col: {ws.max_column}")
for r in range(1, 10):
    vals = [ws.cell(r, c).value for c in range(1, min(16, ws.max_column + 1))]
    print(f"Fila {r}: {vals}")
