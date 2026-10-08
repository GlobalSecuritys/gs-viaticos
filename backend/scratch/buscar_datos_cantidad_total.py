"""Buscar filas no vacias en CANTIDAD TOTAL"""
import openpyxl

ruta = r"c:\gs-viaticos\backend\data\migracion_inventario\CI-FR- INVENTARIO MANTENIMIENTO 2026.xlsx"
wb = openpyxl.load_workbook(ruta, data_only=True)
ws = wb["CANTIDAD TOTAL"]

no_vacias = []
for r in range(1, ws.max_row + 1):
    vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
    if any(v is not None for v in vals):
        no_vacias.append((r, [v for v in vals if v is not None][:6]))
        if len(no_vacias) >= 15:
            break

print(f"Total filas encontradas con datos (primeras 15):")
for r, v in no_vacias:
    print(f"Fila {r}: {v}")
