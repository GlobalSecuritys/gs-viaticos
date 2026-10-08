"""Calcular la tabla dinámica CANTIDAD TOTAL a partir de INVENTARIO GENERAL y de la BD"""
import openpyxl
from collections import defaultdict
import sys, os
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from sqlalchemy import text
from app.database import engine

print("=== 1. DESDE EL EXCEL (Hoja 'INVENTARIO GENERAL') ===")
ruta = r"c:\gs-viaticos\backend\data\migracion_inventario\CI-FR- INVENTARIO MANTENIMIENTO 2026.xlsx"
wb = openpyxl.load_workbook(ruta, data_only=True)
ws = wb["INVENTARIO GENERAL"]

# Encontrar columnas DESCRIPCIÓN y CANTIDAD
header = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
print("Header Excel:", header)

col_desc = None
col_cant = None
for idx, h in enumerate(header, 1):
    if h and "descrip" in str(h).lower():
        col_desc = idx
    if h and "cant" in str(h).lower():
        col_cant = idx

print(f"Columna descripción: {col_desc} ({header[col_desc-1]}), Columna cantidad: {col_cant} ({header[col_cant-1]})")

excel_conteo = defaultdict(float)
excel_filas = 0
for r in range(2, ws.max_row + 1):
    desc = ws.cell(r, col_desc).value
    cant = ws.cell(r, col_cant).value
    if desc:
        desc_clean = str(desc).strip().upper()
        try:
            val_cant = float(cant) if cant is not None else 0.0
        except:
            val_cant = 0.0
        excel_conteo[desc_clean] += val_cant
        excel_filas += 1

print(f"Total filas procesadas en INVENTARIO GENERAL: {excel_filas}")
print(f"Total productos distintos en Excel: {len(excel_conteo)}")
print(f"Suma total de cantidades en Excel: {sum(excel_conteo.values())}")

print("\nTop 15 productos en Excel INVENTARIO GENERAL:")
for p, c in sorted(excel_conteo.items(), key=lambda x: x[1], reverse=True)[:15]:
    print(f"  {p:50} -> {int(c) if c.is_integer() else c}")


print("\n=== 2. DESDE LA BASE DE DATOS (inventario_legacy_items_v2 con cantidad_stock > 0) ===")
with engine.connect() as conn:
    rows = conn.execute(text("""
        SELECT 
            TRIM(UPPER(descripcion)) as producto,
            COUNT(*) as num_registros,
            SUM(cantidad_stock) as suma_stock
        FROM inventario_legacy_items_v2
        WHERE union_temporal::text = 'MANTENIMIENTO'
        GROUP BY TRIM(UPPER(descripcion))
        ORDER BY SUM(cantidad_stock) DESC, TRIM(UPPER(descripcion)) ASC
    """)).mappings().all()

    con_stock = [r for r in rows if r['suma_stock'] > 0]
    sin_stock = [r for r in rows if r['suma_stock'] == 0]

    print(f"Total productos en Legacy: {len(rows)}")
    print(f"Productos con stock > 0: {len(con_stock)}, Suma stock: {sum(r['suma_stock'] for r in con_stock)}")
    print(f"Productos con stock = 0: {len(sin_stock)}")

    print("\nTop 15 productos en BD Legacy con stock:")
    for r in con_stock[:15]:
        print(f"  {r['producto']:50} -> stock: {r['suma_stock']} (en {r['num_registros']} registros)")

    print("\n=== 3. COMPARAR EXCEL vs BD LEGACY ===")
    print(f"Suma Excel INVENTARIO GENERAL: {sum(excel_conteo.values())}")
    print(f"Suma BD Legacy con stock: {sum(r['suma_stock'] for r in con_stock)}")
