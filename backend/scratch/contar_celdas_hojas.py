"""Inspeccionar celdas con datos en todas las hojas"""
import openpyxl

ruta = r"c:\gs-viaticos\backend\data\migracion_inventario\CI-FR- INVENTARIO MANTENIMIENTO 2026.xlsx"
wb = openpyxl.load_workbook(ruta, data_only=True)
for name in wb.sheetnames:
    ws = wb[name]
    celdas = [c.coordinate for row in ws.iter_rows() for c in row if c.value is not None]
    print(f"Hoja '{name}': {len(celdas)} celdas con valor. Rango: {ws.dimensions}")
