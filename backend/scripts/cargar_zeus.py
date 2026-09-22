"""
cargar_zeus.py
==============
Carga / actualiza los ítems del archivo \"EQUIPOS PROYECTO ZEUS.xlsx\" en la tabla
``inventario_items`` con union_temporal = 'PROYECTO_ZEUS'.

Hojas del Excel
---------------
  ODC178  → 11 ítems sin columna de tiempo de entrega.
             estado_entrega = NULL (el color verde del Excel es presentación,
             no hay texto explícito EN STOCK en esta hoja).
  ODC 179 → 14 ítems con columnas Tiempo de Entrega + Observaciones.
             estado_entrega se deriva del texto de la columna:
               «EN STOCK»        → en_stock
               contiene «dia»/«day» (case-insensitive) → en_transito
               «en proceso»      → en_proceso
               cualquier otro texto → en_transito (texto de espera sin formato)

Clave de idempotencia
---------------------
Cada ítem se identifica por ``id_equipo = "zeus:ODC178:N"`` (número de fila de
datos dentro de la hoja). Si el registro ya existe (por corrida anterior del
script viejo), se **actualiza** añadiéndole ``orden_compra`` y
``estado_entrega``. Si no existe, se inserta.

Uso
---
    python scripts/cargar_zeus.py
    python scripts/cargar_zeus.py --dry-run   # solo muestra lo que haría
    python scripts/cargar_zeus.py --reset     # elimina ítems Zeus existentes y re-inserta
"""

import argparse
import sys
import os
import re

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

import openpyxl
from sqlalchemy import select
from app.database import SessionLocal
from app.models.inventario import EstadoEntrega, InventarioItem, UnionTemporal

EXCEL_PATH = os.path.join(
    os.path.dirname(__file__),
    "..",
    "data",
    "migracion_inventario",
    "EQUIPOS PROYECTO ZEUS.xlsx",
)

# Estructura de columnas (índice 0-based dentro de la fila)
# Col 0: vacía; 1: número de artículo; 2: descripción; 3: cantidad;
# 4: tiempo de entrega (solo ODC 179); 5: observaciones (solo ODC 179)
HOJAS = {
    "ODC178": {
        "orden_compra": "ODC178",
        "col_articulo": 1,
        "col_descripcion": 2,
        "col_cantidad": 3,
        "col_tiempo": None,
    },
    "ODC 179": {
        "orden_compra": "ODC179",
        "col_articulo": 1,
        "col_descripcion": 2,
        "col_cantidad": 3,
        "col_tiempo": 4,
    },
}

# Encabezados a ignorar (cualquier celda con este texto en col_articulo)
IGNORAR_ARTICULO = {"número de artículo", "numero de articulo", "n?mero de art?culo"}


def limpiar(val) -> str:
    if val is None:
        return ""
    return str(val).strip()


def derivar_estado_entrega(tiempo_texto: str) -> "EstadoEntrega | None":
    """Convierte el texto libre de tiempo de entrega al enum EstadoEntrega."""
    if not tiempo_texto:
        return None
    t = tiempo_texto.strip().lower()
    if t == "en stock":
        return EstadoEntrega.en_stock
    if "en proceso" in t:
        return EstadoEntrega.en_proceso
    # «8 a 10 Dias», «25 a 30 Dias», «60 a 70 Dias», «8 a 10 days», etc.
    if re.search(r"d[ií]a|day", t, re.IGNORECASE):
        return EstadoEntrega.en_transito
    # Cualquier otro texto con contenido (texto de espera sin formato estándar)
    return EstadoEntrega.en_transito


def cargar(dry_run: bool = False, reset: bool = False):
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    db = SessionLocal()

    insertados = 0
    actualizados = 0
    omitidos = 0
    estado_conteo: dict[str, int] = {}

    try:
        if reset and not dry_run:
            eliminados = (
                db.query(InventarioItem)
                .filter(InventarioItem.union_temporal == UnionTemporal.PROYECTO_ZEUS)
                .delete(synchronize_session=False)
            )
            db.commit()
            print(f"🗑️  Reset: {eliminados} ítems Zeus eliminados.")

        for nombre_hoja, cols in HOJAS.items():
            if nombre_hoja not in wb.sheetnames:
                print(f"⚠️  Hoja '{nombre_hoja}' no encontrada — se omite.")
                continue

            ws = wb[nombre_hoja]
            hoja_clave = cols["orden_compra"]  # "ODC178" o "ODC179"
            orden_compra = cols["orden_compra"]

            fila_n = 0  # contador de filas de datos (excluye encabezados/vacías)
            for row in ws.iter_rows(values_only=True):
                # Saltar filas vacías
                if all(v is None for v in row):
                    continue

                articulo_raw = limpiar(row[cols["col_articulo"]])

                # Saltar encabezados
                if articulo_raw.lower() in IGNORAR_ARTICULO or articulo_raw == "":
                    continue

                fila_n += 1
                clave = f"zeus:{hoja_clave}:{fila_n}"

                numero_articulo = articulo_raw if articulo_raw else None
                descripcion = limpiar(row[cols["col_descripcion"]])[:255]
                try:
                    cantidad = int(row[cols["col_cantidad"]] or 0)
                except (TypeError, ValueError):
                    cantidad = 0

                tiempo_texto = ""
                if cols["col_tiempo"] is not None:
                    tiempo_texto = limpiar(row[cols["col_tiempo"]])

                estado = derivar_estado_entrega(tiempo_texto)
                estado_key = (estado.value if estado else "sin_estado")
                estado_conteo[estado_key] = estado_conteo.get(estado_key, 0) + 1

                print(
                    f"  {'[DRY] ' if dry_run else ''}"
                    f"[{hoja_clave} f{fila_n}] {numero_articulo} — "
                    f"{descripcion[:55]}... ×{cantidad} | "
                    f"entrega={tiempo_texto or '—'} estado={estado_key}"
                )

                if dry_run:
                    insertados += 1
                    continue

                # Buscar si ya existe por id_equipo
                existente = db.scalar(
                    select(InventarioItem).where(InventarioItem.id_equipo == clave)
                )

                if existente:
                    # Actualizar campos Zeus que antes no existían
                    existente.orden_compra = orden_compra
                    existente.estado_entrega = estado
                    # Sincronizar tiempo_entrega también (puede haber cambiado)
                    existente.tiempo_entrega = tiempo_texto[:79] if tiempo_texto else None
                    actualizados += 1
                else:
                    item = InventarioItem(
                        union_temporal=UnionTemporal.PROYECTO_ZEUS,
                        numero_articulo=numero_articulo,
                        descripcion=descripcion,
                        cantidad_stock=cantidad,
                        tiempo_entrega=tiempo_texto[:79] if tiempo_texto else None,
                        orden_compra=orden_compra,
                        estado_entrega=estado,
                        id_equipo=clave,  # clave de idempotencia
                    )
                    db.add(item)
                    insertados += 1

        if not dry_run:
            db.commit()
            print(f"\nCompletado: {insertados} insertados, {actualizados} actualizados, {omitidos} omitidos.")
        else:
            print(f"\nDry-run: {insertados} se procesarian.")

        print("\nDistribucion por estado_entrega:")
        for k, v in sorted(estado_conteo.items()):
            print(f"   {k}: {v} items")

    except Exception as exc:
        db.rollback()
        print(f"\nERROR: {exc}")
        raise
    finally:
        db.close()


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Carga/actualiza EQUIPOS PROYECTO ZEUS")
    parser.add_argument("--dry-run", action="store_true", help="Solo muestra; no hace commit")
    parser.add_argument(
        "--reset",
        action="store_true",
        help="Elimina todos los ítems PROYECTO_ZEUS antes de insertar (requiere confirmación)",
    )
    args = parser.parse_args()

    if args.reset and not args.dry_run:
        resp = input("⚠️  Esto borrará TODOS los ítems Zeus. ¿Continuar? [s/N] ")
        if resp.strip().lower() != "s":
            print("Cancelado.")
            sys.exit(0)

    cargar(dry_run=args.dry_run, reset=args.reset)
