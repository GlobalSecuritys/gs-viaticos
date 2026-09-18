# -*- coding: utf-8 -*-
"""
recargar_tecnicos_exacto.py
===========================
Lee las 6 hojas de técnicos del archivo Excel:
  backend/data/migracion_inventario/CI-FR- INVENTARIO MANTENIMIENTO 2026.xlsx
e inserta con 100% de fidelidad los 77 registros en la tabla
`inventario_tecnicos_items` respetando todas las columnas y fechas reales:
  - fecha_compra
  - factura
  - codigo_barras
  - descripcion
  - cantidad
  - id_equipo
  - serial_gsb
  - concatenado
  - oficina
  - tecnico_nombre
  - fecha_despacho
  - observacion
  - numero_orden
  - oficina_instalada
  - fecha_instalacion
"""

import os
import sys
import datetime
import openpyxl
from sqlalchemy import select, delete

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from app.database import SessionLocal
from app.models.inventario import InventarioTecnicoItem, UnionTemporal
from app.models.usuario import Usuario

EXCEL_PATH = os.path.join(
    os.path.dirname(__file__), "..", "data", "migracion_inventario", "CI-FR- INVENTARIO MANTENIMIENTO 2026.xlsx"
)

# Mapeo de nombre de hoja en Excel -> Usuario en la BD
MAPA_TECNICOS = {
    "GUSTAVO RAMIREZ": {"nombre_db": "Gustavo Ramirez", "id": 37},
    "LUCAS RICO": {"nombre_db": "Lucas Oswaldo Rico Rodriguez", "id": 12},
    "LUIS ANGULO": {"nombre_db": "Luis Francisco Angulo Madera", "id": 13},
    "HECTOR JAVIER CORTES ": {"nombre_db": "HECTOR JAVIER CORTES", "id": 24},
    "WILLIAM RUIZ": {"nombre_db": "William Ruiz", "id": 38},
    "NELSON VARGAS ": {"nombre_db": "Nelson Vargas Calderon", "id": 8},
}


def _limpiar_texto(val):
    if val is None:
        return None
    val_str = str(val).strip()
    if val_str.upper() in {"", "N/A", "NA", "N A", "-", "NONE"}:
        return None
    return val_str


def _formatear_numero_str(val):
    if val is None:
        return None
    if isinstance(val, (int, float)):
        # Si es un número entero exacto, formatear sin .0
        if isinstance(val, float) and val.is_integer():
            return str(int(val))
        return str(val)
    val_str = str(val).strip()
    if val_str.upper() in {"", "N/A", "NA", "N A", "-", "NONE"}:
        return None
    if val_str.endswith(".0"):
        try:
            return str(int(float(val_str)))
        except ValueError:
            pass
    return val_str


def _limpiar_fecha(val):
    if val is None:
        return None
    if isinstance(val, datetime.datetime):
        return val.date()
    if isinstance(val, datetime.date):
        return val
    # Si viene como string
    val_str = str(val).strip()
    if not val_str or val_str.upper() in {"N/A", "NA", "NONE", "-"}:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%Y/%m/%d", "%d-%m-%Y"):
        try:
            return datetime.datetime.strptime(val_str[:10], fmt).date()
        except ValueError:
            continue
    return None


def main():
    print(f"Cargando archivo: {EXCEL_PATH}")
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    db = SessionLocal()

    # Verificar usuarios en la BD
    for sheet_name, info in MAPA_TECNICOS.items():
        u = db.get(Usuario, info["id"])
        if not u:
            # Buscar por nombre
            u = db.query(Usuario).filter(Usuario.nombre.ilike(f"%{info['nombre_db']}%")).first()
            if u:
                info["id"] = u.id
                info["nombre_db"] = u.nombre
            else:
                print(f"ALERTA: Usuario para {sheet_name} no encontrado.")

    print("Limpiando inventario_tecnicos_items actual...")
    db.execute(delete(InventarioTecnicoItem))
    db.commit()

    total_insertados = 0

    for sheet_name, info in MAPA_TECNICOS.items():
        if sheet_name not in wb.sheetnames:
            print(f"Hoja '{sheet_name}' no encontrada en el Excel.")
            continue

        ws = wb[sheet_name]
        headers_raw = next(ws.iter_rows(values_only=True))
        headers = [str(c).strip().replace("\n", " ").upper() if c is not None else "" for c in headers_raw]

        # Índices de columnas
        def get_idx(*nombres):
            for i, h in enumerate(headers):
                for n in nombres:
                    if n.upper() in h:
                        return i
            return -1

        idx_compra = get_idx("FECHA COMPRA", "FECHA DE COMPRA")
        idx_factura = get_idx("FACTURA")
        idx_cod_barras = get_idx("CODIGO DE  BARRAS", "CODIGO DE BARRAS", "BARRAS")
        idx_desc = get_idx("DESCRIPCI")
        idx_cant = get_idx("CANTIDAD")
        idx_id_equipo = get_idx("ID. EQUIPO", "ID EQUIPO")
        idx_serial = get_idx("SERIALES GSB", "SERIAL")
        idx_concat = get_idx("CONCATENADO")
        idx_oficina = get_idx("OFICINA")
        idx_despacho = get_idx("FECHA DE DESPACHO", "DESPACHO")
        idx_obs = get_idx("OBSERVACION")
        idx_orden = get_idx("NUMERO DE ORDEN", "ORDEN")
        idx_ofic_inst = get_idx("OFICINA INSTALDA", "OFICINA INSTALADA")
        idx_inst = get_idx("FECHA DE INSTALACION", "INSTALACION")

        filas = list(ws.iter_rows(values_only=True, min_row=2))
        insertados_hoja = 0

        for r in filas:
            if not any(c is not None for c in r):
                continue

            desc = _limpiar_texto(r[idx_desc]) if idx_desc >= 0 else None
            if not desc:
                continue

            cant_val = r[idx_cant] if idx_cant >= 0 else 1
            try:
                cantidad = int(cant_val) if cant_val else 1
            except (ValueError, TypeError):
                cantidad = 1

            item = InventarioTecnicoItem(
                union_temporal=UnionTemporal.MANTENIMIENTO,
                tecnico_id=info["id"],
                tecnico_nombre=info["nombre_db"],
                descripcion=desc,
                cantidad=max(1, cantidad),
                codigo_barras=_formatear_numero_str(r[idx_cod_barras]) if idx_cod_barras >= 0 else None,
                serial_gsb=_formatear_numero_str(r[idx_serial]) if idx_serial >= 0 else None,
                id_equipo=_formatear_numero_str(r[idx_id_equipo]) if idx_id_equipo >= 0 else None,
                factura=_formatear_numero_str(r[idx_factura]) if idx_factura >= 0 else None,
                fecha_compra=_limpiar_fecha(r[idx_compra]) if idx_compra >= 0 else None,
                oficina=_limpiar_texto(r[idx_oficina]) if idx_oficina >= 0 else None,
                concatenado=_limpiar_texto(r[idx_concat]) if idx_concat >= 0 else None,
                fecha_despacho=_limpiar_fecha(r[idx_despacho]) if idx_despacho >= 0 else None,
                observacion=_limpiar_texto(r[idx_obs]) if idx_obs >= 0 else None,
                numero_orden=_formatear_numero_str(r[idx_orden]) if idx_orden >= 0 else None,
                oficina_instalada=_limpiar_texto(r[idx_ofic_inst]) if idx_ofic_inst >= 0 else None,
                fecha_instalacion=_limpiar_fecha(r[idx_inst]) if idx_inst >= 0 else None,
            )
            db.add(item)
            insertados_hoja += 1

        db.commit()
        print(f"  [OK] Hoja '{sheet_name}' -> Técnico '{info['nombre_db']}': {insertados_hoja} ítems")
        total_insertados += insertados_hoja

    db.close()
    print(f"\nFINALIZADO: Total ítems insertados en inventario_tecnicos_items: {total_insertados}")


if __name__ == "__main__":
    main()
