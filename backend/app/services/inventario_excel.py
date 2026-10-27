"""Servicio para lectura, parseo, búsqueda y caché en memoria de archivos Excel de Inventario.

Los archivos se encuentran en:
- backend/data/migracion_inventario/ (prioridad)
- backend/data/inventario/ (alternativa)

Soporta múltiples archivos y selección automática según la planilla activa.
"""

import datetime
import os
import re
import threading
import unicodedata
import warnings
from typing import Any, Dict, List, Optional, Tuple
import openpyxl

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

# Constante para el piloto Mantenimiento 2026
PLANILLA_PILOTO_EXCEL: int = 1

# Carpetas de búsqueda en orden de prioridad
DIRECTORIOS_DATOS = [
    os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "data", "migracion_inventario")),
    os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "data", "inventario")),
]

_cache_lock = threading.Lock()
# Dict[ruta_absoluta, {"mtime": float, "data": Dict}]
_cache_archivos: Dict[str, Dict[str, Any]] = {}

# Caché en memoria para los cálculos optimizados de Mantenimiento 2026
_cache_mantenimiento_lock = threading.Lock()
_cache_mantenimiento: Dict[str, Dict[str, Any]] = {}


def _normalizar_header(h: Any) -> str:
    """Normaliza encabezados reemplazando saltos de línea por espacio, colapsando espacios y a mayúsculas."""
    if h is None:
        return ""
    s = str(h).replace("\n", " ").replace("\r", " ")
    return re.sub(r"\s+", " ", s).strip().upper()


def _normalizar_nombre_hoja(h: Any) -> str:
    """Normaliza el nombre de una hoja con strip y mayúsculas."""
    if h is None:
        return ""
    return re.sub(r"\s+", " ", str(h)).strip().upper()


def _limpiar_valor(val: Any) -> str:
    """Limpia celdas: floats sin decimales como enteros, 'N/A', 'NA', '-' como vacío."""
    if val is None:
        return ""
    if isinstance(val, float) and val.is_integer():
        return str(int(val))
    s = str(val).strip()
    if s.upper() in ("N/A", "NA", "-", ""):
        return ""
    return s


def _formatear_fecha(val: Any) -> Dict[str, Any]:
    """Formatea una fecha o retorna el texto original con indicador es_invalida si tiene formato anómalo."""
    if val is None:
        return {"valor": "", "es_invalida": False}
    if isinstance(val, (datetime.datetime, datetime.date)):
        if 2020 <= val.year <= 2027:
            return {"valor": val.strftime("%Y-%m-%d"), "es_invalida": False}
        return {"valor": val.strftime("%Y-%m-%d"), "es_invalida": True}

    s = str(val).strip()
    if not s or s.upper() in ("N/A", "NA", "-"):
        return {"valor": "", "es_invalida": False}

    m1 = re.match(r"^(\d{1,2})[/.-](\d{1,2})[/.-](\d{4})$", s)
    if m1:
        d, m, y = int(m1.group(1)), int(m1.group(2)), int(m1.group(3))
        if 2020 <= y <= 2027 and 1 <= m <= 12 and 1 <= d <= 31:
            return {"valor": f"{y:04d}-{m:02d}-{d:02d}", "es_invalida": False}
        return {"valor": s, "es_invalida": True}

    m2 = re.match(r"^(\d{4})[/.-](\d{1,2})[/.-](\d{1,2})$", s)
    if m2:
        y, m, d = int(m2.group(1)), int(m2.group(2)), int(m2.group(3))
        if 2020 <= y <= 2027 and 1 <= m <= 12 and 1 <= d <= 31:
            return {"valor": f"{y:04d}-{m:02d}-{d:02d}", "es_invalida": False}
        return {"valor": s, "es_invalida": True}

    return {"valor": s, "es_invalida": True}


def _formatear_nombre_tecnico(nombre_hoja: str) -> Tuple[str, str]:
    """Retorna el nombre formateado en Title Case y las iniciales."""
    limpio = re.sub(r"\s+", " ", nombre_hoja).strip()
    partes = limpio.split(" ")
    nombre_formateado = " ".join(p.capitalize() for p in partes if p)
    iniciales = "".join(p[0].upper() for p in partes if p)[:2]
    return nombre_formateado, iniciales



def _normalizar_texto(texto: str) -> str:
    """Normaliza texto eliminando tildes, signos y pasando a minúsculas."""
    if not texto:
        return ""
    nfkd = unicodedata.normalize("NFKD", texto)
    sin_tildes = "".join(c for c in nfkd if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]", "", sin_tildes.lower())


def listar_archivos_excel_disponibles() -> List[Dict[str, str]]:
    """Lista todos los archivos .xlsx disponibles en las carpetas de datos."""
    archivos: List[Dict[str, str]] = []
    vistos = set()
    for carpeta in DIRECTORIOS_DATOS:
        if not os.path.exists(carpeta) or not os.path.isdir(carpeta):
            continue
        for nombre in sorted(os.listdir(carpeta)):
            if nombre.lower().endswith(".xlsx") and not nombre.startswith("~$"):
                ruta_completa = os.path.join(carpeta, nombre)
                if os.path.isfile(ruta_completa) and nombre not in vistos:
                    vistos.add(nombre)
                    archivos.append({
                        "nombre": nombre,
                        "ruta": ruta_completa,
                        "carpeta": os.path.basename(carpeta),
                    })
    return archivos


def resolver_ruta_excel(
    planilla_id: Optional[int] = None,
    planilla_nombre: Optional[str] = None,
    archivo: Optional[str] = None,
) -> Tuple[Optional[str], Optional[str]]:
    """Resuelve la ruta absoluta del archivo Excel según el archivo o la planilla indicada.

    Returns:
        (ruta_absoluta, nombre_archivo) o (None, None) si no se encuentra.
    """
    archivos_disponibles = listar_archivos_excel_disponibles()
    if not archivos_disponibles:
        return None, None

    # 1. Si se pidió un archivo explícito
    if archivo and archivo.strip():
        archivo_limpio = os.path.basename(archivo.strip())
        for a in archivos_disponibles:
            if a["nombre"].lower() == archivo_limpio.lower():
                return a["ruta"], a["nombre"]

    # 2. Si se proporcionó planilla_id o el nombre de la planilla
    nombre_pl = (planilla_nombre or "").strip()
    norm_pl = _normalizar_texto(nombre_pl)

    # Regla estricta para Mantenimiento 2026: planilla_id = 1 o mención de mantenimiento
    if planilla_id == PLANILLA_PILOTO_EXCEL or "mantenimiento" in norm_pl:
        matches = [a for a in archivos_disponibles if "mantenimiento" in _normalizar_texto(a["nombre"])]
        if matches:
            # Tomar el de fecha de modificación más reciente
            matches.sort(
                key=lambda a: os.path.getmtime(a["ruta"]) if os.path.exists(a["ruta"]) else 0,
                reverse=True,
            )
            return matches[0]["ruta"], matches[0]["nombre"]

    if norm_pl:
        if "americanglobal" in norm_pl or ("rtc" in norm_pl and "american" in norm_pl):
            for a in archivos_disponibles:
                if "americanglobal" in _normalizar_texto(a["nombre"]):
                    return a["ruta"], a["nombre"]

        if "zeus" in norm_pl:
            for a in archivos_disponibles:
                if "zeus" in _normalizar_texto(a["nombre"]):
                    return a["ruta"], a["nombre"]

        if "comware" in norm_pl:
            for a in archivos_disponibles:
                if "comware" in _normalizar_texto(a["nombre"]):
                    return a["ruta"], a["nombre"]

        if "general" in norm_pl or "globalsecurity" in norm_pl:
            for a in archivos_disponibles:
                if "general" in _normalizar_texto(a["nombre"]):
                    return a["ruta"], a["nombre"]

        # Búsqueda difusa por palabras coincidentes
        palabras = [p for p in re.split(r"\W+", nombre_pl.lower()) if len(p) >= 3]
        for a in archivos_disponibles:
            norm_a = a["nombre"].lower()
            if any(p in norm_a for p in palabras):
                return a["ruta"], a["nombre"]

    # 3. Fallback: primer archivo disponible
    primer = archivos_disponibles[0]
    return primer["ruta"], primer["nombre"]


def _convertir_valor(val: Any) -> Any:
    """Convierte celdas de Excel a cadenas seguras para serializar en JSON."""
    if val is None:
        return ""
    if hasattr(val, "isoformat"):
        try:
            return val.strftime("%Y-%m-%d") if hasattr(val, "strftime") else val.isoformat()
        except Exception:
            return str(val)
    if isinstance(val, float):
        if val.is_integer():
            return str(int(val))
        return f"{val:.2f}".rstrip("0").rstrip(".")
    if isinstance(val, int):
        return str(val)
    s = str(val).strip()
    if s.upper() in ("N/A", "NA", "-", ""):
        return ""
    return s


def _cargar_excel_desde_disco(ruta_archivo: str, nombre_archivo: str) -> Dict[str, Any]:
    """Carga todas las hojas del archivo Excel con sus columnas limpias y filas."""
    wb = openpyxl.load_workbook(ruta_archivo, data_only=True, read_only=True)
    hojas_raw = wb.sheetnames

    # Excluir hojas rotas conocidas en Mantenimiento
    es_mantenimiento = "mantenimiento" in _normalizar_texto(nombre_archivo)
    if es_mantenimiento:
        hojas = [h for h in hojas_raw if _normalizar_nombre_hoja(h) not in {"CANTIDAD TOTAL", "FILTRO ACTA"}]
    else:
        hojas = hojas_raw

    datos_hojas: Dict[str, Dict[str, Any]] = {}

    for sheet_name in hojas:
        ws = wb[sheet_name]
        filas_iter = ws.iter_rows(values_only=True)

        encabezados: List[str] = []
        filas: List[Dict[str, Any]] = []
        cols_mapeo: List[Tuple[int, str]] = []
        desc_col_idx: Optional[int] = None

        for row in filas_iter:
            if not row or all(c is None for c in row):
                continue

            # Primera fila no vacía como encabezados
            if not encabezados:
                encabezados_raw = [
                    re.sub(r"\s+", " ", str(c).replace("\n", " ").replace("\r", " ")).strip()
                    if c is not None and str(c).strip()
                    else f"Col_{i+1}"
                    for i, c in enumerate(row)
                ]

                # Filtrar CONCATENADO en Mantenimiento
                cols_filtradas = []
                for i, col in enumerate(encabezados_raw):
                    if es_mantenimiento and col.upper().startswith("CONCATENADO"):
                        continue
                    cols_filtradas.append((i, col))

                # Asegurar columnas únicas
                nombres_vistos: Dict[str, int] = {}
                cols_unicas: List[Tuple[int, str]] = []
                for orig_idx, col in cols_filtradas:
                    if col in nombres_vistos:
                        nombres_vistos[col] += 1
                        c_name = f"{col}_{nombres_vistos[col]}"
                    else:
                        nombres_vistos[col] = 1
                        c_name = col
                    cols_unicas.append((orig_idx, c_name))
                    if _normalizar_header(c_name) in {"DESCRIPCIÓN", "DESCRIPCION"}:
                        desc_col_idx = orig_idx

                encabezados = [c[1] for c in cols_unicas]
                cols_mapeo = cols_unicas
                continue

            # Descartar filas fantasma en INVENTARIO GENERAL sin DESCRIPCIÓN
            if es_mantenimiento and _normalizar_nombre_hoja(sheet_name) == "INVENTARIO GENERAL":
                desc_val = row[desc_col_idx] if desc_col_idx is not None and desc_col_idx < len(row) else None
                if not _limpiar_valor(desc_val):
                    continue

            fila_dict: Dict[str, Any] = {}
            tiene_datos = False
            for orig_idx, h in cols_mapeo:
                val = row[orig_idx] if orig_idx < len(row) else None
                convertido = _convertir_valor(val)
                fila_dict[h] = convertido
                if convertido:
                    tiene_datos = True

            if tiene_datos:
                filas.append(fila_dict)

        datos_hojas[sheet_name] = {
            "columnas": encabezados,
            "total_filas": len(filas),
            "filas": filas,
        }

    wb.close()
    return {
        "existe": True,
        "archivo": nombre_archivo,
        "ruta": ruta_archivo,
        "hojas": hojas,
        "datos_hojas": datos_hojas,
    }


def _obtener_datos_archivo_cacheados(ruta: str, nombre: str, forzar_recarga: bool = False) -> Dict[str, Any]:
    global _cache_archivos
    mtime_actual = os.path.getmtime(ruta)

    with _cache_lock:
        cached = _cache_archivos.get(ruta)
        if forzar_recarga or not cached or cached.get("mtime") != mtime_actual:
            data = _cargar_excel_desde_disco(ruta, nombre)
            _cache_archivos[ruta] = {"mtime": mtime_actual, "data": data}
            return data
        return cached["data"]


def obtener_datos_inventario_excel(
    planilla_id: Optional[int] = None,
    planilla_nombre: Optional[str] = None,
    archivo: Optional[str] = None,
    hoja: Optional[str] = None,
    limit: int = 1500,
    offset: int = 0,
    forzar_recarga: bool = False,
) -> Dict[str, Any]:
    """Retorna los datos del Excel correspondiente a la planilla o archivo indicado."""
    ruta, nombre_archivo = resolver_ruta_excel(
        planilla_id=planilla_id,
        planilla_nombre=planilla_nombre,
        archivo=archivo,
    )

    if not ruta or not os.path.exists(ruta):
        ref_nombre = archivo or planilla_nombre or (f"Planilla {planilla_id}" if planilla_id else "inventario")
        return {
            "existe": False,
            "archivo": ref_nombre if ref_nombre.endswith(".xlsx") else f"{ref_nombre}.xlsx",
            "ruta": "",
            "hojas": [],
            "hoja_activa": None,
            "columnas": [],
            "total_filas": 0,
            "filas": [],
            "mensaje": f"No se encontró un archivo Excel para '{ref_nombre}' en backend/data/migracion_inventario/.",
        }

    datos_completos = _obtener_datos_archivo_cacheados(ruta, nombre_archivo, forzar_recarga=forzar_recarga)

    hojas = datos_completos.get("hojas", [])
    datos_hojas = datos_completos.get("datos_hojas", {})

    # Determinar la hoja activa
    hoja_activa = None
    if hoja and hoja in hojas:
        hoja_activa = hoja
    elif "INVENTARIO GENERAL" in hojas and datos_hojas.get("INVENTARIO GENERAL", {}).get("total_filas", 0) > 0:
        hoja_activa = "INVENTARIO GENERAL"
    else:
        # Primera hoja que tenga filas
        for h in hojas:
            if datos_hojas.get(h, {}).get("total_filas", 0) > 0:
                hoja_activa = h
                break
        if not hoja_activa and hojas:
            hoja_activa = hojas[0]

    hoja_info = datos_hojas.get(hoja_activa, {}) if hoja_activa else {}
    columnas = hoja_info.get("columnas", [])
    todas_las_filas = hoja_info.get("filas", [])
    total_filas = len(todas_las_filas)

    # Paginación
    inicio = max(0, offset)
    fin = inicio + limit
    filas_paginadas = todas_las_filas[inicio:fin]

    return {
        "existe": True,
        "archivo": nombre_archivo,
        "hojas": hojas,
        "hoja_activa": hoja_activa,
        "columnas": columnas,
        "total_filas": total_filas,
        "filas": filas_paginadas,
        "mensaje": None,
    }


def buscar_en_excel(
    hoja: Optional[str] = None,
    serial: Optional[str] = None,
    oficina: Optional[str] = None,
    tecnico: Optional[str] = None,
    fecha: Optional[str] = None,
    planilla_id: Optional[int] = None,
    planilla_nombre: Optional[str] = None,
    archivo: Optional[str] = None,
) -> Dict[str, Any]:
    """Busca registros en el archivo Excel según criterios progresivos."""
    ruta, nombre_archivo = resolver_ruta_excel(
        planilla_id=planilla_id,
        planilla_nombre=planilla_nombre,
        archivo=archivo,
    )

    criterios = {
        "serial": (serial or "").strip(),
        "oficina": (oficina or "").strip(),
        "tecnico": (tecnico or "").strip(),
        "fecha": (fecha or "").strip(),
    }

    if not ruta or not os.path.exists(ruta):
        return {
            "total_encontrados": 0,
            "hoja": hoja or "",
            "criterios": criterios,
            "resultados": [],
        }

    datos_completos = _obtener_datos_archivo_cacheados(ruta, nombre_archivo)
    datos_hojas = datos_completos.get("datos_hojas", {})
    hojas = datos_completos.get("hojas", [])

    hojas_a_buscar = [hoja] if (hoja and hoja in datos_hojas) else hojas

    resultados: List[Dict[str, Any]] = []

    def coincide_criterio(fila: Dict[str, Any], valor_buscado: str, posibles_cols: List[str]) -> bool:
        if not valor_buscado:
            return True
        v_low = valor_buscado.lower()

        # 1. Probar en columnas sugeridas
        for col_name, val in fila.items():
            c_low = col_name.lower()
            if any(p in c_low for p in posibles_cols):
                if v_low in str(val or "").lower():
                    return True

        # 2. Si no encontró en columnas específicas, buscar en toda la fila
        return any(v_low in str(val or "").lower() for val in fila.values())

    for h_name in hojas_a_buscar:
        info = datos_hojas.get(h_name, {})
        for fila in info.get("filas", []):
            if not coincide_criterio(fila, criterios["serial"], ["serial", "codigo", "código", "id", "sds"]):
                continue
            if not coincide_criterio(fila, criterios["oficina"], ["oficina", "sede", "destino", "cliente"]):
                continue
            if not coincide_criterio(fila, criterios["tecnico"], ["tecnico", "técnico", "responsable", "nombre"]):
                continue
            if not coincide_criterio(fila, criterios["fecha"], ["fecha"]):
                continue

            fila_con_hoja = dict(fila)
            fila_con_hoja["_hoja_origen"] = h_name
            resultados.append(fila_con_hoja)
            if len(resultados) >= 500:
                break
        if len(resultados) >= 500:
            break

    return {
        "total_encontrados": len(resultados),
        "hoja": hoja or (hojas[0] if hojas else ""),
        "criterios": criterios,
        "resultados": resultados,
    }


# =============================================================================
# MÓDULO OPTIMIZADO EXCLUSIVO PARA MANTENIMIENTO 2026 (PLANILLA_PILOTO_EXCEL = 1)
# =============================================================================

def _calcular_datos_mantenimiento(ruta_archivo: str, nombre_archivo: str) -> Dict[str, Any]:
    """Realiza todos los cálculos determinísticos sobre el Excel de Mantenimiento 2026."""
    wb = openpyxl.load_workbook(ruta_archivo, data_only=True, read_only=True)
    mtime = os.path.getmtime(ruta_archivo)
    fecha_act = datetime.datetime.fromtimestamp(mtime).strftime("%d/%m/%Y %H:%M")

    # 1. INVENTARIO GENERAL
    ws_ig = wb["INVENTARIO GENERAL"]
    rows_ig = list(ws_ig.iter_rows(values_only=True))
    h_ig = [_normalizar_header(h) for h in rows_ig[0]]
    d_idx = h_ig.index("DESCRIPCIÓN") if "DESCRIPCIÓN" in h_ig else h_ig.index("DESCRIPCION")
    s_idx = h_ig.index("SERIALES GSB")

    ig_agrupado: Dict[str, Dict[str, Any]] = {}
    total_ig_validas = 0
    total_con_serial = 0
    total_sin_serial = 0

    for r in rows_ig[1:]:
        desc_clean = _limpiar_valor(r[d_idx])
        if not desc_clean:
            continue
        total_ig_validas += 1
        serial = _limpiar_valor(r[s_idx])
        if serial:
            total_con_serial += 1
        else:
            total_sin_serial += 1

        k = _normalizar_header(desc_clean)
        if k not in ig_agrupado:
            ig_agrupado[k] = {
                "descripcion": desc_clean,
                "bodega": 0,
                "con_serial": 0,
                "sin_serial": 0,
            }
        ig_agrupado[k]["bodega"] += 1
        if serial:
            ig_agrupado[k]["con_serial"] += 1
        else:
            ig_agrupado[k]["sin_serial"] += 1

    # 2. SALIDAS
    ws_sal = wb["SALIDAS "]
    rows_sal = list(ws_sal.iter_rows(values_only=True))
    h_sal = [_normalizar_header(h) for h in rows_sal[0]]
    d_sal_idx = h_sal.index("DESCRIPCIÓN") if "DESCRIPCIÓN" in h_sal else h_sal.index("DESCRIPCION")
    s_sal_idx = h_sal.index("SERIALES GSB")
    ofic_sal_idx = [i for i, h in enumerate(h_sal) if "OFICINA INSTAL" in h][0]
    fecha_inst_sal_idx = [i for i, h in enumerate(h_sal) if "FECHA DE INSTALACION" in h][0]
    tecnico_sal_idx = [i for i, h in enumerate(h_sal) if "TECNICO" in h][0]
    ofic_orig_idx = [i for i, h in enumerate(h_sal) if h == "OFICINA"][0]
    fecha_desp_idx = [i for i, h in enumerate(h_sal) if "FECHA DE DESPACHO" in h][0]
    obs_idx = [i for i, h in enumerate(h_sal) if "OBSERVACION" in h][0] if any("OBSERVACION" in h for h in h_sal) else None
    orden_idx = [i for i, h in enumerate(h_sal) if "ORDEN" in h][0] if any("ORDEN" in h for h in h_sal) else None

    salidas_agrupado: Dict[str, Dict[str, int]] = {}
    salidas_filas: List[Dict[str, Any]] = []
    conteo_estados = {"instalada": 0, "pendiente": 0, "devuelta": 0, "danada": 0}
    tecnicos_salidas_set = set()

    for idx, r in enumerate(rows_sal[1:]):
        desc_clean = _limpiar_valor(r[d_sal_idx])
        if not desc_clean:
            continue
        ofic_inst = _limpiar_valor(r[ofic_sal_idx])
        fecha_inst_val = r[fecha_inst_sal_idx]
        ofic_u = ofic_inst.upper()

        if "DAÑAD" in ofic_u or "DANAD" in ofic_u:
            estado = "danada"
        elif "BODEGA" in ofic_u or "DEVUELT" in ofic_u:
            estado = "devuelta"
        elif ofic_inst and _limpiar_valor(fecha_inst_val):
            estado = "instalada"
        else:
            estado = "pendiente"

        conteo_estados[estado] += 1
        tecnico_val = _limpiar_valor(r[tecnico_sal_idx])
        if tecnico_val:
            tecnicos_salidas_set.add(tecnico_val)

        k = _normalizar_header(desc_clean)
        if k not in salidas_agrupado:
            salidas_agrupado[k] = {"pendientes": 0, "instaladas": 0}
        if estado == "pendiente":
            salidas_agrupado[k]["pendientes"] += 1
        elif estado == "instalada":
            salidas_agrupado[k]["instaladas"] += 1

        salidas_filas.append({
            "id": idx + 1,
            "descripcion": desc_clean,
            "serial": _limpiar_valor(r[s_sal_idx]),
            "tecnico": tecnico_val,
            "oficina": _limpiar_valor(r[ofic_orig_idx]),
            "fecha_despacho": _formatear_fecha(r[fecha_desp_idx]),
            "oficina_instalada": ofic_inst,
            "fecha_instalacion": _formatear_fecha(fecha_inst_val),
            "estado": estado,
            "observacion": _limpiar_valor(r[obs_idx]) if obs_idx is not None else "",
            "orden": _limpiar_valor(r[orden_idx]) if orden_idx is not None else "",
        })

    # 3. CANTIDAD TOTAL (Unión de productos de INVENTARIO GENERAL y SALIDAS)
    todos_keys = sorted(list(set(ig_agrupado.keys()).union(set(salidas_agrupado.keys()))))
    cantidad_total_filas: List[Dict[str, Any]] = []
    prods_con_stock = 0
    prods_sin_stock = 0

    for k in todos_keys:
        ig_d = ig_agrupado.get(k, {"descripcion": k, "bodega": 0, "con_serial": 0, "sin_serial": 0})
        sal_d = salidas_agrupado.get(k, {"pendientes": 0, "instaladas": 0})
        desc = ig_d.get("descripcion") or k
        bodega = ig_d["bodega"]
        pend = sal_d["pendientes"]
        inst = sal_d["instaladas"]

        if bodega > 0:
            prods_con_stock += 1
        else:
            prods_sin_stock += 1

        sin_stock_con_salidas = (bodega == 0 and (pend > 0 or inst > 0))

        cantidad_total_filas.append({
            "descripcion": desc,
            "unidades_bodega": bodega,
            "con_serial": ig_d["con_serial"],
            "sin_serial": ig_d["sin_serial"],
            "salidas_pendientes": pend,
            "instaladas": inst,
            "sin_stock_con_salidas": sin_stock_con_salidas,
        })

    cantidad_total_filas.sort(key=lambda x: (x["unidades_bodega"], x["salidas_pendientes"]), reverse=True)

    # 4. PRESTAMOS (8 filas válidas, ignorando bloque suelto D:E desde fila 33)
    ws_pres = wb["PRESTAMOS"]
    rows_pres = list(ws_pres.iter_rows(values_only=True))
    prestamos_filas = []
    total_prestamos_unidades = 0
    for r in rows_pres[1:]:
        desc = _limpiar_valor(r[0])
        cant = _limpiar_valor(r[1])
        if desc and cant:
            try:
                cnt_num = int(float(cant))
                prestamos_filas.append({"descripcion": desc, "cantidad": cnt_num})
                total_prestamos_unidades += cnt_num
            except Exception:
                pass

    # 5. COMPRAS (205 filas válidas)
    ws_com = wb["COMPRAS "]
    rows_com = list(ws_com.iter_rows(values_only=True))
    h_com = [_normalizar_header(h) for h in rows_com[0]]
    d_com_idx = h_com.index("DESCRIPCIÓN") if "DESCRIPCIÓN" in h_com else h_com.index("DESCRIPCION")
    f_com_idx = [i for i, h in enumerate(h_com) if "FECHA" in h][0]
    fac_com_idx = [i for i, h in enumerate(h_com) if "FACTURA" in h][0]
    cb_com_idx = [i for i, h in enumerate(h_com) if "BARRAS" in h][0]
    cant_com_idx = [i for i, h in enumerate(h_com) if "CANTIDAD" in h][0]
    sds_com_idx = [i for i, h in enumerate(h_com) if "SDS" in h][0]
    ser_com_idx = [i for i, h in enumerate(h_com) if "SERIAL" in h][0]

    compras_filas = []
    for idx, r in enumerate(rows_com[1:]):
        desc = _limpiar_valor(r[d_com_idx])
        if not desc:
            continue
        compras_filas.append({
            "id": idx + 1,
            "fecha_compra": _formatear_fecha(r[f_com_idx]),
            "factura": _limpiar_valor(r[fac_com_idx]),
            "codigo_barras": _limpiar_valor(r[cb_com_idx]),
            "descripcion": desc,
            "cantidad": 1,
            "no_sds": _limpiar_valor(r[sds_com_idx]),
            "serial": _limpiar_valor(r[ser_com_idx]),
        })

    # 6. TECNICOS ACTIVOS (Detectados dinámicamente)
    HOJAS_RESERVADAS = {
        "INVENTARIO GENERAL", "SALIDAS", "COMPRAS", "PRESTAMOS",
        "CANTIDAD TOTAL", "FILTRO ACTA"
    }
    tecnicos_lista = []
    detalle_tecnicos: Dict[str, Dict[str, Any]] = {}
    total_unidades_tecnicos = 0

    for sheet in wb.sheetnames:
        norm_s = _normalizar_nombre_hoja(sheet)
        if norm_s in HOJAS_RESERVADAS:
            continue
        ws_t = wb[sheet]
        rows_t = list(ws_t.iter_rows(values_only=True))
        if not rows_t:
            continue
        h_t = [_normalizar_header(h) for h in rows_t[0]]
        if any("TECNICO" in h for h in h_t) and any("SERIALES GSB" in h for h in h_t):
            d_idx_t = h_t.index("DESCRIPCIÓN") if "DESCRIPCIÓN" in h_t else (h_t.index("DESCRIPCION") if "DESCRIPCION" in h_t else None)
            if d_idx_t is None:
                continue

            s_idx_t = [i for i, h in enumerate(h_t) if "SERIAL" in h][0]
            f_desp_idx_t = [i for i, h in enumerate(h_t) if "FECHA DE DESPACHO" in h][0]
            ofic_idx_t = [i for i, h in enumerate(h_t) if h == "OFICINA"][0]
            ofic_inst_idx_t = [i for i, h in enumerate(h_t) if "OFICINA INSTAL" in h][0]
            f_inst_idx_t = [i for i, h in enumerate(h_t) if "FECHA DE INSTALACION" in h][0]
            obs_idx_t = [i for i, h in enumerate(h_t) if "OBSERVACION" in h][0] if any("OBSERVACION" in h for h in h_t) else None
            orden_idx_t = [i for i, h in enumerate(h_t) if "ORDEN" in h][0] if any("ORDEN" in h for h in h_t) else None

            filas_t = []
            for t_idx, r in enumerate(rows_t[1:]):
                desc_t = _limpiar_valor(r[d_idx_t])
                if not desc_t:
                    continue
                filas_t.append({
                    "id": t_idx + 1,
                    "descripcion": desc_t,
                    "serial": _limpiar_valor(r[s_idx_t]),
                    "fecha_despacho": _formatear_fecha(r[f_desp_idx_t]),
                    "oficina": _limpiar_valor(r[ofic_idx_t]),
                    "oficina_instalada": _limpiar_valor(r[ofic_inst_idx_t]),
                    "fecha_instalacion": _formatear_fecha(r[f_inst_idx_t]),
                    "observacion": _limpiar_valor(r[obs_idx_t]) if obs_idx_t is not None else "",
                    "orden": _limpiar_valor(r[orden_idx_t]) if orden_idx_t is not None else "",
                })

            if filas_t:
                nom_formateado, ini = _formatear_nombre_tecnico(sheet)
                t_info = {
                    "hoja": sheet.strip(),
                    "nombre": nom_formateado,
                    "iniciales": ini,
                    "total_unidades": len(filas_t),
                }
                tecnicos_lista.append(t_info)
                total_unidades_tecnicos += len(filas_t)
                detalle_tecnicos[norm_s] = {
                    "hoja": sheet.strip(),
                    "nombre": nom_formateado,
                    "iniciales": ini,
                    "total_unidades": len(filas_t),
                    "items": filas_t,
                }

    wb.close()

    resumen = {
        "indicadores": {
            "unidades_bodega": total_ig_validas,
            "salidas_totales": len(salidas_filas),
            "pendientes_instalar": conteo_estados["pendiente"],
            "fecha_actualizacion_excel": fecha_act,
        },
        "resumen_botones": {
            "cantidad_total": {
                "unidades": total_ig_validas,
                "productos": len(cantidad_total_filas),
                "con_stock": prods_con_stock,
                "sin_stock": prods_sin_stock,
            },
            "salidas": {
                "total": len(salidas_filas),
                "pendientes": conteo_estados["pendiente"],
                "instaladas": conteo_estados["instalada"],
                "devueltas": conteo_estados["devuelta"],
                "danadas": conteo_estados["danada"],
            },
            "prestamos": {
                "total_unidades": total_prestamos_unidades,
                "lineas": len(prestamos_filas),
            },
            "prestamos_tecnicos": {
                "total_tecnicos": len(tecnicos_lista),
                "total_unidades": total_unidades_tecnicos,
            },
            "compras": {
                "total_unidades": len(compras_filas),
            },
        },
    }

    resumen_totales_cantidad_total = {
        "total_unidades_bodega": total_ig_validas,
        "total_con_serial": total_con_serial,
        "total_sin_serial": total_sin_serial,
        "total_salidas_pendientes": conteo_estados["pendiente"],
        "total_instaladas": conteo_estados["instalada"],
        "total_productos": len(cantidad_total_filas),
        "productos_con_stock_bodega": prods_con_stock,
        "productos_sin_stock_bodega": prods_sin_stock,
    }

    return {
        "resumen": resumen,
        "cantidad_total": {
            "resumen_totales": resumen_totales_cantidad_total,
            "items": cantidad_total_filas,
        },
        "salidas": {
            "resumen_estados": conteo_estados,
            "tecnicos_disponibles": sorted(list(tecnicos_salidas_set)),
            "items": salidas_filas,
        },
        "prestamos": {
            "total_unidades": total_prestamos_unidades,
            "items": prestamos_filas,
        },
        "compras": {
            "total_unidades": len(compras_filas),
            "items": compras_filas,
        },
        "tecnicos": tecnicos_lista,
        "detalle_tecnicos": detalle_tecnicos,
    }


def _obtener_calculados_mantenimiento(ruta_archivo: str, nombre_archivo: str, forzar_recarga: bool = False) -> Dict[str, Any]:
    """Retorna los datos calculados desde caché o recalcula si el mtime ha cambiado."""
    global _cache_mantenimiento
    mtime = os.path.getmtime(ruta_archivo)

    with _cache_mantenimiento_lock:
        cached = _cache_mantenimiento.get(ruta_archivo)
        if forzar_recarga or not cached or cached.get("mtime") != mtime:
            datos = _calcular_datos_mantenimiento(ruta_archivo, nombre_archivo)
            _cache_mantenimiento[ruta_archivo] = {"mtime": mtime, "datos": datos}
            return datos
        return cached["datos"]


def _resolver_ruta_mantenimiento(planilla_id: int = PLANILLA_PILOTO_EXCEL) -> Tuple[str, str]:
    from fastapi import HTTPException, status
    if planilla_id != PLANILLA_PILOTO_EXCEL:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Este módulo optimizado está disponible únicamente para Mantenimiento 2026 (planilla {PLANILLA_PILOTO_EXCEL}).",
        )
    ruta, nombre = resolver_ruta_excel(planilla_id=planilla_id, planilla_nombre="Mantenimiento 2026")
    if not ruta or not os.path.exists(ruta):
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="No se encontró el archivo Excel de Mantenimiento 2026 en el servidor.",
        )
    return ruta, nombre


def obtener_resumen_mantenimiento(planilla_id: int = PLANILLA_PILOTO_EXCEL) -> Dict[str, Any]:
    """Retorna los indicadores de encabezado y conteos para las tarjetas resumen."""
    ruta, nombre = _resolver_ruta_mantenimiento(planilla_id)
    datos = _obtener_calculados_mantenimiento(ruta, nombre)
    return datos["resumen"]


def obtener_vista_mantenimiento(
    planilla_id: int,
    vista: str,
    q: Optional[str] = None,
    filtro_estado: Optional[str] = None,
    filtro_tecnico: Optional[str] = None,
    page: int = 1,
    page_size: int = 50,
) -> Dict[str, Any]:
    """Retorna los datos filtrados y paginados de una de las 4 vistas."""
    from fastapi import HTTPException, status
    ruta, nombre = _resolver_ruta_mantenimiento(planilla_id)
    datos = _obtener_calculados_mantenimiento(ruta, nombre)

    vista_norm = (vista or "").strip().lower()
    q_norm = (q or "").strip().lower()

    if vista_norm == "cantidad_total":
        seccion = datos["cantidad_total"]
        items = seccion["items"]
        if q_norm:
            items = [it for it in items if q_norm in it["descripcion"].lower()]
        total = len(items)
        inicio = max(0, (page - 1) * page_size)
        paginados = items[inicio : inicio + page_size]
        total_paginas = max(1, (total + page_size - 1) // page_size)
        return {
            "vista": "cantidad_total",
            "total": total,
            "page": page,
            "page_size": page_size,
            "total_paginas": total_paginas,
            "resumen_totales": seccion["resumen_totales"],
            "items": paginados,
        }

    elif vista_norm == "salidas":
        seccion = datos["salidas"]
        items = seccion["items"]
        if q_norm:
            items = [
                it for it in items
                if (
                    q_norm in it["descripcion"].lower()
                    or q_norm in it["serial"].lower()
                    or q_norm in it["tecnico"].lower()
                    or q_norm in it["oficina"].lower()
                    or q_norm in it["oficina_instalada"].lower()
                )
            ]
        if filtro_estado and filtro_estado.strip().lower() not in ("todas", "todos", ""):
            f_est = filtro_estado.strip().lower()
            items = [it for it in items if it["estado"] == f_est]
        if filtro_tecnico and filtro_tecnico.strip().lower() not in ("todos", ""):
            f_tec = filtro_tecnico.strip().lower()
            items = [it for it in items if f_tec in it["tecnico"].lower()]

        total = len(items)
        inicio = max(0, (page - 1) * page_size)
        paginados = items[inicio : inicio + page_size]
        total_paginas = max(1, (total + page_size - 1) // page_size)
        return {
            "vista": "salidas",
            "total": total,
            "page": page,
            "page_size": page_size,
            "total_paginas": total_paginas,
            "resumen_estados": seccion["resumen_estados"],
            "tecnicos_disponibles": seccion["tecnicos_disponibles"],
            "items": paginados,
        }

    elif vista_norm == "prestamos":
        seccion = datos["prestamos"]
        items = seccion["items"]
        if q_norm:
            items = [it for it in items if q_norm in it["descripcion"].lower()]
        total = len(items)
        inicio = max(0, (page - 1) * page_size)
        paginados = items[inicio : inicio + page_size]
        total_paginas = max(1, (total + page_size - 1) // page_size)
        return {
            "vista": "prestamos",
            "total": total,
            "total_unidades": seccion["total_unidades"],
            "page": page,
            "page_size": page_size,
            "total_paginas": total_paginas,
            "items": paginados,
        }

    elif vista_norm == "compras":
        seccion = datos["compras"]
        items = seccion["items"]
        if q_norm:
            items = [
                it for it in items
                if (
                    q_norm in it["descripcion"].lower()
                    or q_norm in it["serial"].lower()
                    or q_norm in it["factura"].lower()
                    or q_norm in it["codigo_barras"].lower()
                    or q_norm in it["no_sds"].lower()
                )
            ]
        total = len(items)
        inicio = max(0, (page - 1) * page_size)
        paginados = items[inicio : inicio + page_size]
        total_paginas = max(1, (total + page_size - 1) // page_size)
        return {
            "vista": "compras",
            "total": total,
            "total_unidades": seccion["total_unidades"],
            "page": page,
            "page_size": page_size,
            "total_paginas": total_paginas,
            "items": paginados,
        }

    else:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Vista desconocida: '{vista}'. Opciones válidas: cantidad_total, salidas, prestamos, compras.",
        )


def obtener_tecnicos_mantenimiento(planilla_id: int = PLANILLA_PILOTO_EXCEL) -> List[Dict[str, Any]]:
    """Retorna la lista liviana de técnicos activos con hojas válidas para el slider."""
    ruta, nombre = _resolver_ruta_mantenimiento(planilla_id)
    datos = _obtener_calculados_mantenimiento(ruta, nombre)
    return datos["tecnicos"]


def obtener_detalle_tecnico_mantenimiento(
    planilla_id: int,
    nombre_hoja: str,
    q: Optional[str] = None,
    page: int = 1,
    page_size: int = 50,
) -> Dict[str, Any]:
    """Retorna el detalle de préstamos de un técnico específico. Rechaza con 404 si no es técnico activo."""
    from fastapi import HTTPException, status
    ruta, nombre = _resolver_ruta_mantenimiento(planilla_id)
    datos = _obtener_calculados_mantenimiento(ruta, nombre)

    norm_buscada = _normalizar_nombre_hoja(nombre_hoja)
    detalle_map = datos["detalle_tecnicos"]

    if norm_buscada not in detalle_map:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="La hoja solicitada no corresponde a un técnico activo con préstamos.",
        )

    tecnico_info = detalle_map[norm_buscada]
    items = tecnico_info["items"]
    q_norm = (q or "").strip().lower()
    if q_norm:
        items = [
            it for it in items
            if (
                q_norm in it["descripcion"].lower()
                or q_norm in it["serial"].lower()
                or q_norm in it["oficina"].lower()
                or q_norm in it["oficina_instalada"].lower()
            )
        ]

    total = len(items)
    inicio = max(0, (page - 1) * page_size)
    paginados = items[inicio : inicio + page_size]
    total_paginas = max(1, (total + page_size - 1) // page_size)

    return {
        "hoja": tecnico_info["hoja"],
        "nombre": tecnico_info["nombre"],
        "iniciales": tecnico_info["iniciales"],
        "total_unidades": tecnico_info["total_unidades"],
        "total": total,
        "page": page,
        "page_size": page_size,
        "total_paginas": total_paginas,
        "items": paginados,
    }